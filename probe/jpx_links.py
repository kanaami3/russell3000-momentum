import io
from curl_cffi import requests
import openpyxl
s = requests.Session(impersonate="chrome")
B = "https://www.jpx.co.jp"
for u in ["/markets/statistics-equities/investor-type/t13vrt000002gv8k-att/stock_1_w_20261005_20261009.xlsx",
          "/markets/statistics-equities/margin/tvdivq0000001rk9-att/20261009_mtcurrent.xlsx"]:
    r = s.get(B + u, timeout=30)
    print("==", u, r.status_code, len(r.content))
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    for ws in wb.worksheets:
        print("--- SHEET", ws.title, ws.max_row, ws.max_column)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 70: break
            vals = [("" if v is None else str(v)) for v in row]
            while vals and vals[-1] == "": vals.pop()
            if vals: print(i, "|".join(vals)[:260])
